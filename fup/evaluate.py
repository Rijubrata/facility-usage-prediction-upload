"""Evaluate the saved model on the untouched chronological holdout (Nov-Dec 2025).

Outputs (in ``reports/``):
  prediction_review.xlsx   Summary (model vs baselines) / Predictions / Error analysis / Facility confusion
  prediction_review.html   self-contained, filterable UI table of the same predictions
  test_predictions.csv     flat file with every prediction vs actual
  metrics.json             all metrics for the model and baselines

Run:  python -m fup.evaluate
"""
from __future__ import annotations

import json

import joblib
import numpy as np
import pandas as pd

from . import config as C
from . import features as Fe
from . import model as M
from .html_report import write_html
from .metrics import evaluate
from .report import review_table
from .split import split_states

METRIC_LABELS = {
    "facility_acc": ("Facility", "Accuracy (top-1)"),
    "facility_top3": ("Facility", "Top-3 accuracy"),
    "facility_macro_f1": ("Facility", "Macro F1 (imbalance-aware)"),
    "facility_logloss": ("Facility", "Log loss (lower is better)"),
    "day_acc": ("Usage day", "Exact calendar day"),
    "day_within1": ("Usage day", "Within +/-1 day"),
    "day_mae": ("Usage day", "MAE (days)"),
    "hour_acc": ("Usage hour", "Exact hour"),
    "hour_within1": ("Usage hour", "Within +/-1 hour"),
    "hour_mae": ("Usage hour", "MAE (hours)"),
    "nudge_within60": ("Notification time", f"Within +/-{C.NUDGE_TOLERANCE_MIN} min of actual booking"),
    "nudge_within180": ("Notification time", "Within +/-3 h of actual booking"),
    "nudge_median_ae_h": ("Notification time", "Median abs. error (hours)"),
    "nudge_mae_h": ("Notification time", "MAE (hours)"),
    "nudge_before_booking": ("Notification time", "Nudge sent at/before actual booking"),
    "mean_output_accuracy": ("Overall", "Mean output accuracy (avg of 4 match rates)"),
    "mean_outputs_matched": ("Overall", "Mean outputs matched per record (of 4)"),
    "all4_match": ("Overall", "All 4 outputs match"),
    "fac_day_hour_match": ("Overall", "Facility + day + hour all match"),
    "usage_slot_within2h": ("Overall", "Right facility & usage time within 2 h"),
}


def _segment_table(rev: pd.DataFrame, col: str, name: str) -> pd.DataFrame:
    g = rev.groupby(col, observed=True)
    t = pd.DataFrame({
        "segment": name,
        "value": g.size().index.astype(str),
        "n": g.size().to_numpy(),
        "facility_acc": g.facility_match.mean().to_numpy(),
        "day_acc": g.day_match.mean().to_numpy(),
        "hour_acc": g.hour_match.mean().to_numpy(),
        "nudge_acc": g.nudge_match.mean().to_numpy(),
        "mean_output_acc": (g.outputs_matched.mean() / 4).to_numpy(),
        "all4": (g.outputs_matched.apply(lambda x: (x == 4).mean())).to_numpy(),
    })
    return t


def error_analysis(rev: pd.DataFrame, s: Fe.State) -> pd.DataFrame:
    rev = rev.copy()
    rev["history_bucket"] = pd.cut(rev.history_n, [0, 5, 20, 50, 100, 10_000],
                                   labels=["1-5", "6-20", "21-50", "51-100", "100+"])
    rev["pred_conf_bucket"] = pd.cut(rev.confidence_facility, [0, 0.4, 0.6, 0.8, 0.9, 1.0],
                                     labels=["<0.4", "0.4-0.6", "0.6-0.8", "0.8-0.9", "0.9+"])
    lead_d = (s.meta.y_use.to_numpy() // Fe.DAY) - (s.meta.y_book.to_numpy() // Fe.DAY)
    rev["actual_lead"] = pd.cut(lead_d, [-1, 0, 1, 3, 10_000], labels=["same day", "1 day", "2-3 days", "4+ days"])
    rev["actual_dow"] = pd.Categorical(rev.act_day.str[:3], C.DAYS)
    rev["switched_facility"] = np.where(s.base.last_fac.to_numpy() == s.meta.y_fac.to_numpy(),
                                        "same as previous booking", "different from previous")
    parts = [_segment_table(rev, c, n) for c, n in [
        ("history_bucket", "Bookings in history"), ("act_facility", "Actual facility"),
        ("actual_lead", "Actual booking lead"), ("actual_dow", "Actual usage weekday"),
        ("switched_facility", "Facility vs previous booking"), ("pred_conf_bucket", "Facility confidence")]]

    # generator ground truth -- used ONLY here to explain errors, never as a model input
    if C.TRUTH_CSV.exists():
        truth = pd.read_csv(C.TRUTH_CSV, dtype=str).fillna("")
        rev = rev.merge(truth, left_on="resident", right_on="resident_id", how="left")
        drift = pd.to_datetime(rev.drift_date.replace("", None))
        rev["drift_status"] = np.select(
            [drift.isna(), drift < C.VAL_START, drift < C.TEST_START],
            ["no drift", "drift before Sep (seen in training)", "drift in Sep-Oct (little post-drift history)"],
            "drift during test (unseen)")
        rev["drift_kind"] = rev.drift_type.replace("", "none")
        parts += [_segment_table(rev, c, n) for c, n in [
            ("persona", "Resident persona (generator)"), ("lead_type", "Booking style (generator)"),
            ("drift_status", "Behaviour drift (generator)"), ("drift_kind", "Drift type (generator)")]]
    return pd.concat(parts, ignore_index=True)


def summary_frame(metrics: dict, baselines: dict) -> pd.DataFrame:
    rows = []
    for k, (out, label) in METRIC_LABELS.items():
        row = {"output": out, "metric": label, "model": metrics[k]}
        for name, bm in baselines.items():
            row[name] = bm[k]
        rows.append(row)
    return pd.DataFrame(rows)


def main(argv=None) -> None:
    model: M.ChainModel = joblib.load(C.MODEL_DIR / "chain_model.joblib")
    df = Fe.load_bookings()
    sp = split_states(Fe.build_state(df))
    te = sp["test"]
    print(f"holdout: {len(te.meta):,} target bookings created {C.TEST_START.date()} .. {C.DATA_END.date()} (exclusive)")

    p = model.predict(te)
    metrics = evaluate(te, p)
    baselines = {}
    for name, b in [("baseline_habit", M.HeuristicBaseline()), ("baseline_popularity", M.PopularityBaseline())]:
        baselines[name] = evaluate(te, b.fit(sp["train_val"]).predict(te))

    rev = review_table(df, te, p)
    summ = summary_frame(metrics, baselines)
    err = error_analysis(rev, te)
    labels = [C.FACILITY_LABEL[f] for f in C.FACILITY_IDS]
    conf = pd.crosstab(pd.Categorical(rev.act_facility, labels), pd.Categorical(rev.pred_facility, labels),
                       rownames=["actual"], colnames=["predicted"], dropna=False)

    C.REPORT_DIR.mkdir(exist_ok=True)
    rev.to_csv(C.REPORT_DIR / "test_predictions.csv", index=False)
    (C.REPORT_DIR / "metrics.json").write_text(json.dumps(
        {"model": metrics, **baselines}, indent=2, default=lambda x: round(float(x), 4)))
    write_excel(rev, summ, err, conf)
    write_html(rev, summ, err, C.REPORT_DIR / "prediction_review.html")

    print(summ.assign(**{c: summ[c].map(lambda v: f"{v:.3f}") for c in ["model", *baselines]}).to_string(index=False))
    print(f"\nwrote reports -> {C.REPORT_DIR}")


def write_excel(rev, summ, err, conf) -> None:
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Font, PatternFill

    path = C.REPORT_DIR / "prediction_review.xlsx"
    green, red = PatternFill("solid", fgColor="D9F2E1"), PatternFill("solid", fgColor="F8D7DA")
    head = PatternFill("solid", fgColor="EFE3F1")
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        intro = pd.DataFrame({"Facility Usage Prediction - holdout review": [
            f"Holdout = every booking created {C.TEST_START.date()} to {(C.DATA_END - pd.Timedelta(days=1)).date()} "
            f"({len(rev):,} records). Never used for fitting or tuning.",
            "Each prediction is made at the moment the resident created their previous booking ('predicted_at') "
            "using only data available then.",
            f"Match rules: facility exact; day = exact calendar date; hour = exact start hour; "
            f"nudge = within +/-{C.NUDGE_TOLERANCE_MIN} min of the actual booking-creation time.",
            "Columns 'baseline_*' are simple non-ML references evaluated on the same records.", ""]})
        intro.to_excel(xw, sheet_name="Summary", index=False)
        summ.to_excel(xw, sheet_name="Summary", index=False, startrow=len(intro) + 2)
        rev.to_excel(xw, sheet_name="Predictions", index=False)
        err.to_excel(xw, sheet_name="Error analysis", index=False)
        conf.to_excel(xw, sheet_name="Facility confusion")

        wb = xw.book
        ws = wb["Summary"]
        ws.column_dimensions["A"].width = 20
        ws.column_dimensions["B"].width = 48
        for c in "CDE":
            ws.column_dimensions[c].width = 20
        ws["A1"].font = Font(bold=True, size=13)
        for row in ws.iter_rows(min_row=len(intro) + 4, min_col=3, max_col=5):
            for cell in row:
                cell.number_format = "0.000"
        for cell in ws[len(intro) + 3]:
            cell.font, cell.fill = Font(bold=True), head

        ws = wb["Predictions"]
        widths = dict(A=11, B=8, C=17, D=9, E=52, F=13, G=12, H=9, I=17, J=13, K=12, L=9, M=17)
        for col, w in widths.items():
            ws.column_dimensions[col].width = w
        for cell in ws[1]:
            cell.font, cell.fill = Font(bold=True), head
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
            row[0].alignment = Alignment(wrap_text=True, vertical="top")
        last = ws.max_row
        for col in "NOPQ":  # boolean match flags
            rng = f"{col}2:{col}{last}"
            ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=["TRUE"], fill=green))
            ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=["FALSE"], fill=red))
        ws.conditional_formatting.add(f"V2:V{last}", CellIsRule(operator="equal", formula=['"YES"'], fill=green))
        ws.conditional_formatting.add(f"V2:V{last}", CellIsRule(operator="equal", formula=['"NO"'], fill=red))
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions

        ws = wb["Error analysis"]
        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 42
        for row in ws.iter_rows(min_row=2, min_col=4, max_col=9):
            for cell in row:
                cell.number_format = "0.0%"
        for cell in ws[1]:
            cell.font, cell.fill = Font(bold=True), head
        ws.auto_filter.ref = ws.dimensions
    print(f"wrote {path}")


if __name__ == "__main__":
    main()
